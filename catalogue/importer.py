"""
Validated catalogue import from CSV / Excel (.xlsx).

Flow (used by both the owner web page and the `import_products` command):
    1. read_table()      - file -> list of raw row dicts (no database access)
    2. build_preview()   - validate every row, detect duplicates, compare with the DB
    3. apply_rows()      - one database transaction; any exception rolls everything back

Stock quantities are intentionally NOT imported (stock source not verified).
Nothing here touches the shop's existing billing database; it only reads an export file.
"""
import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import IntegrityError, transaction

from . import units as U
from .models import Category, Product, ProductAlias
from .textnorm import loose_key

MAX_ROWS = 20000
MAX_BYTES = 5 * 1024 * 1024

COLUMN_SYNONYMS = {
    "external_code": {"external_code", "code", "product_code", "item_code", "sku", "id"},
    "name_en": {"name_en", "name", "english_name", "product", "item", "product_name", "item_name"},
    "name_ta": {"name_ta", "tamil_name", "tamil"},
    "category": {"category", "group", "department"},
    "price": {"price", "rate", "selling_price", "sale_price"},
    "unit": {"unit", "uom"},
    "pack_description": {"pack_description", "pack", "pack_size"},
    "min_quantity": {"min_quantity", "min_qty", "minimum", "min"},
    "quantity_step": {"quantity_step", "step", "increment"},
    "is_available": {"is_available", "available"},
    "aliases": {"aliases", "alias", "other_names"},
}

UNIT_WORDS = {
    "piece": U.PIECE, "pieces": U.PIECE, "pc": U.PIECE, "pcs": U.PIECE, "nos": U.PIECE, "no": U.PIECE, "each": U.PIECE,
    "packet": U.PACKET, "packets": U.PACKET, "pkt": U.PACKET,
    "bottle": U.BOTTLE, "bottles": U.BOTTLE, "btl": U.BOTTLE,
    "pack": U.PACK, "packs": U.PACK,
    "kg": U.KG, "kgs": U.KG, "kilo": U.KG, "kilogram": U.KG,
    "g": U.GRAM, "gm": U.GRAM, "gms": U.GRAM, "gram": U.GRAM, "grams": U.GRAM,
    "l": U.LITRE, "lt": U.LITRE, "ltr": U.LITRE, "litre": U.LITRE, "liter": U.LITRE, "litres": U.LITRE,
    "ml": U.ML,
}
for _code in U.UNIT_INFO:  # also accept the canonical codes themselves
    UNIT_WORDS[_code.lower()] = _code

TRUE_WORDS = {"1", "y", "yes", "true", "available", "in stock"}
FALSE_WORDS = {"0", "n", "no", "false", "unavailable", "out of stock"}


class ImportFileError(Exception):
    """The file itself cannot be used (wrong type, unreadable, too big...)."""


# ---------------------------------------------------------------- reading

def _canonical_header(h):
    key = (h or "").strip().lower().replace(" ", "_").replace("-", "_")
    for canon, names in COLUMN_SYNONYMS.items():
        if key in names:
            return canon
    return None


def read_table(fileobj, filename):
    data = fileobj.read()
    if len(data) > MAX_BYTES:
        raise ImportFileError("File is larger than 5 MB.")
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        raw_rows = _read_csv(data)
    elif name.endswith(".xlsx"):
        raw_rows = _read_xlsx(data)
    else:
        raise ImportFileError("Please upload a .csv or .xlsx file.")
    if not raw_rows:
        raise ImportFileError("The file is empty.")
    header = [_canonical_header(h) for h in raw_rows[0]]
    if "name_en" not in header or "price" not in header:
        raise ImportFileError("The first row must contain at least 'name' and 'price' columns. "
                              "See docs/IMPORT_FORMAT.md for the expected columns.")
    rows = []
    for i, cells in enumerate(raw_rows[1:], start=2):
        if not any((c or "").strip() for c in cells):
            continue
        row = {"_line": i}
        for col, val in zip(header, cells):
            if col and col not in row:
                row[col] = (val or "").strip()
        rows.append(row)
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Too many rows (max {MAX_ROWS}).")
    if not rows:
        raise ImportFileError("No product rows found under the header.")
    return rows, [h for h in header if h]


def _read_csv(data):
    for enc in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(enc)
            break
        except UnicodeError:
            continue
    else:
        raise ImportFileError("Could not read the file. Please save it as 'CSV UTF-8' so Tamil text is kept.")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[c for c in r] for r in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportFileError("Excel import needs the 'openpyxl' package (see requirements.txt).") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError("Could not open the Excel file.") from exc
    ws = wb.worksheets[0]
    out = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if v is None else (format(v, "f") if isinstance(v, Decimal) else str(v)) for v in row])
        if len(out) > MAX_ROWS + 1:
            break
    wb.close()
    return out


# ---------------------------------------------------------------- validating

def _dec(text, field, errors, places=None):
    try:
        d = Decimal(text.replace(",", ""))
    except (InvalidOperation, AttributeError):
        errors.append(f"{field}: '{text}' is not a number")
        return None
    if not d.is_finite() or d <= 0:
        errors.append(f"{field}: must be greater than zero")
        return None
    if places is not None and d != d.quantize(Decimal(1).scaleb(-places)):
        errors.append(f"{field}: at most {places} decimal places allowed")
        return None
    return d


def build_preview(rows, columns):
    """Validate rows. Returns (clean_rows, summary). Nothing is written to the database."""
    existing_by_code = {p.external_code: p for p in Product.objects.exclude(external_code=None)}
    existing_by_name = {(p.name_en.lower(), p.pack_description.lower()): p for p in Product.objects.all()}
    seen_codes, seen_names = {}, {}
    clean, counts = [], {"create": 0, "update": 0, "unchanged": 0, "error": 0}
    price_changes = 0

    for row in rows:
        errors, warnings = [], []
        out = {"line": row["_line"], "errors": errors, "warnings": warnings}
        name_en = row.get("name_en", "")
        if not name_en:
            errors.append("name is missing")
        elif len(name_en) > 200:
            errors.append("name is longer than 200 characters")
        code = row.get("external_code", "") or None
        pack = row.get("pack_description", "")
        out.update(name_en=name_en, name_ta=row.get("name_ta", ""), external_code=code, pack_description=pack)

        price = _dec(row.get("price", ""), "price", errors, places=2) if row.get("price", "") != "" else None
        if row.get("price", "") == "":
            errors.append("price is missing")
        out["price"] = str(price) if price is not None else ""

        unit_raw = row.get("unit", "").strip().lower()
        unit = UNIT_WORDS.get(unit_raw)
        if "unit" not in columns or not unit_raw:
            errors.append("unit is missing (needed so customers are never billed the wrong way)")
        elif unit is None:
            errors.append(f"unit '{row.get('unit')}' is not recognised")
        out["unit"] = unit or ""

        for field in ("min_quantity", "quantity_step"):
            raw = row.get(field, "")
            if raw == "":
                out[field] = "1"
                if unit and U.is_measure(unit):
                    warnings.append(f"{field} not given for a measured product; defaulted to 1 - please check")
            else:
                val = _dec(raw, field, errors, places=3)
                out[field] = str(val) if val is not None else ""
        if not errors and Decimal(out["min_quantity"]) % Decimal(out["quantity_step"]) != 0:
            errors.append("min_quantity must be a multiple of quantity_step")
        if not errors and unit and not U.is_measure(unit) and Decimal(out["quantity_step"]) % 1 != 0:
            errors.append("counted items need a whole-number quantity_step")

        cat = row.get("category", "")
        if not cat:
            cat = "Uncategorised"
            warnings.append("category missing; will use 'Uncategorised'")
        out["category"] = cat

        avail = row.get("is_available", "").strip().lower()
        if avail == "":
            out["is_available"] = None
        elif avail in TRUE_WORDS:
            out["is_available"] = True
        elif avail in FALSE_WORDS:
            out["is_available"] = False
        else:
            errors.append(f"available: '{row.get('is_available')}' not understood (use yes/no)")
            out["is_available"] = None

        aliases = [a.strip() for a in row.get("aliases", "").split("|") if a.strip()]
        out["aliases"] = aliases

        # duplicates inside the file
        if code:
            if code in seen_codes:
                errors.append(f"duplicate product code '{code}' (also on line {seen_codes[code]})")
            seen_codes[code] = seen_codes.get(code, row["_line"])
        if name_en:
            nk = (name_en.lower(), pack.lower())
            if nk in seen_names:
                errors.append(f"duplicate product name '{name_en}' (also on line {seen_names[nk]}); "
                              f"add a pack description to tell them apart")
            seen_names[nk] = seen_names.get(nk, row["_line"])

        # compare with the database
        existing = existing_by_code.get(code) if code else None
        if existing is None and name_en:
            existing = existing_by_name.get((name_en.lower(), pack.lower()))
            if existing is not None and code and existing.external_code and existing.external_code != code:
                errors.append(f"name already belongs to product with code '{existing.external_code}'")
                existing = None
        if errors:
            out["action"] = "error"
            counts["error"] += 1
        elif existing is None:
            out["action"] = "create"
            counts["create"] += 1
        else:
            out["existing_id"] = existing.pk
            changed = (existing.price != Decimal(out["price"]) or existing.unit != out["unit"]
                       or existing.name_ta != (out["name_ta"] or existing.name_ta)
                       or existing.min_quantity != Decimal(out["min_quantity"])
                       or existing.quantity_step != Decimal(out["quantity_step"])
                       or (out["is_available"] is not None and existing.is_available != out["is_available"])
                       or bool(aliases))
            if existing.price != Decimal(out["price"]):
                out["old_price"] = str(existing.price)
                price_changes += 1
            out["action"] = "update" if changed else "unchanged"
            counts[out["action"]] += 1
        clean.append(out)

    summary = {"counts": counts, "total": len(clean), "price_changes": price_changes}
    return clean, summary


# ---------------------------------------------------------------- applying

def apply_rows(rows, user=None):
    """
    Apply previewed rows atomically. Rows with errors are skipped.
    Returns a report dict. Any unexpected exception rolls back the whole import.
    """
    report = {"created": 0, "updated": 0, "skipped_errors": 0, "alias_warnings": []}
    with transaction.atomic():
        for r in rows:
            if r["action"] in ("error", "unchanged"):
                if r["action"] == "error":
                    report["skipped_errors"] += 1
                continue
            cat = Category.objects.filter(name_en__iexact=r["category"]).first() \
                or Category.objects.create(name_en=r["category"])
            if r["action"] == "create":
                product = Product(category=cat, external_code=r["external_code"], name_en=r["name_en"],
                                  name_ta=r["name_ta"], price=Decimal(r["price"]), unit=r["unit"],
                                  pack_description=r["pack_description"],
                                  min_quantity=Decimal(r["min_quantity"]), quantity_step=Decimal(r["quantity_step"]),
                                  is_available=True if r["is_available"] is None else r["is_available"])
                product._changed_by = user
                product.full_clean()
                product.save()
                report["created"] += 1
            else:
                product = Product.objects.select_for_update().get(pk=r["existing_id"])
                product._changed_by = user
                product.price = Decimal(r["price"])
                product.unit = r["unit"]
                product.min_quantity = Decimal(r["min_quantity"])
                product.quantity_step = Decimal(r["quantity_step"])
                if r["name_ta"]:
                    product.name_ta = r["name_ta"]
                if r["external_code"] and not product.external_code:
                    product.external_code = r["external_code"]
                if r["is_available"] is not None:
                    product.is_available = r["is_available"]
                product.category = cat
                product.full_clean()
                product.save()
                report["updated"] += 1
            for alias in r.get("aliases", []):
                key = loose_key(alias)
                if not key:
                    continue
                if ProductAlias.objects.filter(normalized=key).exists():
                    report["alias_warnings"].append(f"Alias '{alias}' skipped (already in use)")
                    continue
                a = ProductAlias(product=product, text=alias)
                try:
                    a.full_clean()
                except Exception as exc:
                    report["alias_warnings"].append(f"Alias '{alias}' skipped ({exc})")
                    continue
                try:
                    with transaction.atomic():
                        a.save()
                except IntegrityError:
                    report["alias_warnings"].append(f"Alias '{alias}' skipped (duplicate)")
    return report
